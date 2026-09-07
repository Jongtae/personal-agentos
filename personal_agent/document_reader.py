"""Local, bounded extraction for documents inside a user-connected folder."""
from dataclasses import dataclass
from pathlib import Path
import zipfile

MAX_FILE_BYTES=10_000_000
MAX_EXTRACTED_CHARS=120_000
MAX_ZIP_MEMBERS=2_000
MAX_ZIP_UNCOMPRESSED=30_000_000
SUPPORTED_SUFFIXES={'.txt':'text','.md':'markdown','.pdf':'pdf','.docx':'docx','.xlsx':'xlsx'}


@dataclass
class Document:
    kind:str
    segments:list

    @property
    def text(self):
        return '\n'.join(segment['text'] for segment in self.segments)[:MAX_EXTRACTED_CHARS]


def supported(path):
    return Path(path).suffix.lower() in SUPPORTED_SUFFIXES


def _bounded_zip(path):
    try:
        with zipfile.ZipFile(path) as archive:
            entries=archive.infolist()
            total=sum(entry.file_size for entry in entries)
            if len(entries)>MAX_ZIP_MEMBERS or total>MAX_ZIP_UNCOMPRESSED:
                raise ValueError('문서 압축 해제 크기가 제한을 초과합니다.')
            if any(entry.file_size and entry.compress_size and entry.file_size/entry.compress_size>200 for entry in entries):
                raise ValueError('안전하지 않은 압축 문서입니다.')
    except zipfile.BadZipFile as exc:
        raise ValueError('손상되었거나 지원하지 않는 Office 문서입니다.') from exc


def _text(path):
    try: raw=path.read_bytes()
    except OSError as exc: raise ValueError('문서를 읽을 수 없습니다.') from exc
    if b'\0' in raw: raise ValueError('일반 텍스트 문서가 아닙니다.')
    try: value=raw.decode('utf-8')
    except UnicodeDecodeError as exc: raise ValueError('UTF-8 텍스트 문서가 아닙니다.') from exc
    return Document(SUPPORTED_SUFFIXES[path.suffix.lower()],[{'location':f'줄 {index}','text':line} for index,line in enumerate(value.splitlines(),1) if line][:MAX_EXTRACTED_CHARS])


def _pdf(path):
    try:
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted: raise ValueError('암호화된 PDF는 읽을 수 없습니다.')
        segments=[]
        for page_number,page in enumerate(reader.pages,1):
            value=page.extract_text() or ''
            if value.strip(): segments.append({'location':f'페이지 {page_number}','text':value.strip()})
            if sum(len(s['text']) for s in segments)>=MAX_EXTRACTED_CHARS: break
        if not segments: raise ValueError('텍스트를 추출할 수 없는 PDF입니다. 이미지 OCR은 아직 지원하지 않습니다.')
        return Document('pdf',segments)
    except ValueError: raise
    except Exception as exc: raise ValueError('PDF를 읽을 수 없습니다.') from exc


def _docx(path):
    _bounded_zip(path)
    try:
        from docx import Document as WordDocument
        document=WordDocument(path);segments=[]
        for index,paragraph in enumerate(document.paragraphs,1):
            if paragraph.text.strip():segments.append({'location':f'문단 {index}','text':paragraph.text.strip()})
        for table_number,table in enumerate(document.tables,1):
            for row_number,row in enumerate(table.rows,1):
                value=' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if value:segments.append({'location':f'표 {table_number} · 행 {row_number}','text':value})
        if not segments:raise ValueError('텍스트가 없는 DOCX 문서입니다.')
        return Document('docx',segments)
    except ValueError:raise
    except Exception as exc:raise ValueError('DOCX 문서를 읽을 수 없습니다.') from exc


def _column(number):
    result=''
    while number:
        number,remainder=divmod(number-1,26);result=chr(65+remainder)+result
    return result


def _xlsx(path):
    _bounded_zip(path)
    try:
        from openpyxl import load_workbook
        workbook=load_workbook(path,read_only=True,data_only=True);segments=[]
        for sheet in workbook.worksheets:
            for row_number,row in enumerate(sheet.iter_rows(values_only=True),1):
                cells=[(index,value) for index,value in enumerate(row,1) if value not in (None,'')]
                if cells:
                    start,end=cells[0][0],cells[-1][0]
                    value=' | '.join(str(cell) for _,cell in cells)
                    segments.append({'location':f'{sheet.title}!{_column(start)}{row_number}:{_column(end)}{row_number}','text':value})
                if len(segments)>=10_000 or sum(len(s['text']) for s in segments)>=MAX_EXTRACTED_CHARS:break
        if not segments:raise ValueError('읽을 수 있는 값이 없는 XLSX 문서입니다.')
        return Document('xlsx',segments)
    except ValueError:raise
    except Exception as exc:raise ValueError('XLSX 문서를 읽을 수 없습니다.') from exc


def read(path):
    path=Path(path)
    if not path.is_file() or path.stat().st_size>MAX_FILE_BYTES:raise ValueError('10MB 이하의 지원 문서만 읽을 수 있습니다.')
    kind=SUPPORTED_SUFFIXES.get(path.suffix.lower())
    if not kind:raise ValueError('지원 형식은 TXT, MD, PDF, DOCX, XLSX입니다.')
    if kind in ('text','markdown'):return _text(path)
    if kind=='pdf':return _pdf(path)
    if kind=='docx':return _docx(path)
    return _xlsx(path)
