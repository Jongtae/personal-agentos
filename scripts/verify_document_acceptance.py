"""Live P2-01 acceptance against an isolated connected-document folder."""
import json
import os
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from personal_agent.quickstart_store import QuickStore
from personal_agent.quickstart_service import AgentService


def make_pdf(path):
 writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
 font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
 page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
 stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 50 200 Td (PDF Orion launch date 2032-04-20) Tj ET');page[NameObject('/Contents')]=stream
 with path.open('wb') as output:writer.write(output)


report_path=Path(os.environ.get('AGENTOS_DOCUMENT_ACCEPTANCE_REPORT',str(Path(tempfile.gettempdir())/'agentos-document-acceptance.json')))
source=QuickStore(Path.home()/'.local/share/agentos')
with tempfile.TemporaryDirectory() as root:
 docs=Path(root)/'documents';docs.mkdir()
 (docs/'orion.txt').write_text('TXT Orion owner: Mina.')
 (docs/'orion.md').write_text('# Orion\nMD Orion budget: 7300 USD')
 make_pdf(docs/'orion.pdf')
 word=WordDocument();word.add_paragraph('DOCX Orion status: approved');word.save(docs/'orion.docx')
 book=Workbook();sheet=book.active;sheet.title='Orion';sheet.append(['Milestone','Date']);sheet.append(['Launch','2032-04-20']);book.save(docs/'orion.xlsx')
 store=QuickStore(Path(root)/'data');service=AgentService(store);service.save_roots({'paths':[str(docs)]})
 config=dict(source.config('model'));checked=source.config('model_test',{})
 if config.get('model')=='openrouter/free' and checked.get('runtime_model'):config['model']=checked['runtime_model']
 service.save_model({**config,'api_key':source.secret('model_key')})
 checked=service.test_model();results=[]
 if checked['ok']:
  for label,prompt,filename,location in [
   ('txt','orion.txt에서 소유자를 찾아줘. 파일과 줄 근거를 표시해줘.','orion.txt','줄'),
   ('markdown','orion.md에서 예산을 찾아줘. 파일과 줄 근거를 표시해줘.','orion.md','줄'),
   ('pdf','orion.pdf에서 출시일을 찾아줘. 파일과 페이지 근거를 표시해줘.','orion.pdf','페이지'),
   ('docx','orion.docx에서 상태를 찾아줘. 파일과 문단 근거를 표시해줘.','orion.docx','문단'),
   ('xlsx','orion.xlsx에서 출시일을 찾아줘. 파일과 시트 및 셀 근거를 표시해줘.','orion.xlsx','Orion!')]:
   identifier=store.enqueue(prompt,'document-'+label);service.run_one();job=next(row for row in store.jobs() if row['id']==identifier)
   text=job.get('response') or ''
   passed=job['status'] in ('succeeded','partial') and filename in text and location in text
   results.append({'case':label,'passed':passed,'status':job['status'],'model':job['model'],'response':text[:500]})
 report={'tested_at':time.time(),'model_validation':checked,'results':results}
 report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2))
 print(json.dumps([(row['case'],row['passed']) for row in results],ensure_ascii=False))
 if not checked['ok'] or not all(row['passed'] for row in results):raise SystemExit(1)
