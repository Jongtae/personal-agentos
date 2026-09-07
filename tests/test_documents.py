import tempfile
import unittest
from pathlib import Path
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from docx import Document as WordDocument
from openpyxl import Workbook
from personal_agent.agent_runtime import Capabilities
from personal_agent.document_reader import read
from personal_agent.quickstart_store import QuickStore


class DocumentTests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)/'connected';self.root.mkdir()
  (self.root/'brief.txt').write_text('Aurora launch date is 2031-10-12.\nOwner Mina.')
  (self.root/'notes.md').write_text('# Aurora\nBudget: 4200 USD')
  word=WordDocument();word.add_paragraph('DOCX Aurora owner Mina');table=word.add_table(rows=1,cols=2);table.cell(0,0).text='Status';table.cell(0,1).text='Approved';word.save(self.root/'report.docx')
  book=Workbook();sheet=book.active;sheet.title='Launch';sheet.append(['Project','Date']);sheet.append(['Aurora','2031-10-12']);book.save(self.root/'plan.xlsx')
  writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
  font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
  page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
  stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 50 200 Td (PDF Aurora budget 4200) Tj ET');page[NameObject('/Contents')]=stream
  with (self.root/'brief.pdf').open('wb') as output:writer.write(output)
  self.store=QuickStore(Path(self.temp.name)/'data');self.store.put('file_roots',[{'id':'root','path':str(self.root)}]);self.caps=Capabilities(self.store,None,{},'','job',lambda *args:None)
 def tearDown(self):self.temp.cleanup()
 def test_extracts_supported_types_with_locations(self):
  expected={'brief.txt':('Aurora','줄 1'),'notes.md':('Budget','줄 2'),'brief.pdf':('budget','페이지 1'),'report.docx':('owner','문단 1'),'plan.xlsx':('2031-10-12','Launch!A2:B2')}
  for name,(text,location) in expected.items():
   result=read(self.root/name)
   self.assertIn(text,result.text);self.assertIn(location,[s['location'] for s in result.segments])
 def test_capability_returns_safe_locations_and_searches_documents(self):
  result=self.caps.read_file('root','plan.xlsx')
  self.assertEqual(result['sources'],['파일: plan.xlsx · Launch!A1:B1'])
  self.assertIn('[Launch!A2:B2]',result['content'])
  hits=self.caps.find_files('Aurora')['files']
  self.assertEqual({hit['path'] for hit in hits},{'brief.txt','notes.md','brief.pdf','report.docx','plan.xlsx'})
  self.assertTrue(all(hit['location'] for hit in hits))
 def test_rejects_non_document_and_path_escape(self):
  (self.root/'binary.bin').write_bytes(b'\x00bad')
  with self.assertRaises(ValueError):self.caps.read_file('root','binary.bin')
  with self.assertRaises(ValueError):self.caps.read_file('root','../data/private/quickstart.db')
 def test_rejects_bad_office_document(self):
  (self.root/'bad.docx').write_text('not a zip')
  with self.assertRaises(ValueError):self.caps.read_file('root','bad.docx')


if __name__=='__main__':unittest.main()
